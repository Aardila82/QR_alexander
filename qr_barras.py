from PIL import Image, ImageDraw, ImageFont
import qrcode
import barcode
from barcode.writer import ImageWriter
from docx import Document
from docx.shared import Inches, Pt
from docx.shared import Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os


# --- Configuración ---
dpi = 300  # Resolución para impresión

# Medidas en centímetros
lado_qr_cm = 1        # QR pequeño (1x1 cm)
ancho_barcode_cm = 5  # ancho del código de barras
alto_barcode_cm = 0.5 # alto del código de barras

# Conversión a píxeles
lado_qr_px = int(lado_qr_cm / 2.54 * dpi)
ancho_barcode_px = int(ancho_barcode_cm / 2.54 * dpi)
alto_barcode_px = int(alto_barcode_cm / 2.54 * dpi)

# --- 1. Generar código QR ---
qr_data = "https://example.com"
qr = qrcode.QRCode(box_size=10, border=2)
qr.add_data(qr_data)
qr.make(fit=True)
qr_img = qr.make_image(fill_color="black", back_color="white").convert("RGBA")
qr_img = qr_img.resize((lado_qr_px, lado_qr_px))

# --- 2. Generar código de barras SIN NÚMEROS ABAJO ---
codigo_valor = "123456789012"
ean = barcode.get('code128', codigo_valor, writer=ImageWriter())
barcode_img_path = "barcode.png"

if os.path.exists(barcode_img_path):
    os.remove(barcode_img_path)

# Guardar código de barras con fondo blanco
ean.save("barcode", {
    "module_height": max(1, alto_barcode_px / 4),
    "module_width": 0.10,  # barras más finas
    "font_size": 4,
    "dpi": dpi,
    "text_distance": 0,
    "write_text": False
})

# Abrir y convertir blanco a transparente
barcode_img = Image.open(barcode_img_path).convert("RGBA")
datas = barcode_img.getdata()
nueva_data = []
for item in datas:
    # Detectar el blanco puro y volverlo transparente
    if item[0] > 240 and item[1] > 240 and item[2] > 240:
        nueva_data.append((255, 255, 255, 0))
    else:
        nueva_data.append(item)
barcode_img.putdata(nueva_data)

barcode_img = barcode_img.resize((ancho_barcode_px, alto_barcode_px))

# --- 3. Dibujar los números ARRIBA ---
font_size = 10
try:
    font = ImageFont.truetype("arial.ttf", font_size)
except Exception:
    font = ImageFont.load_default()

temp = Image.new("RGB", (10, 10))
draw_temp = ImageDraw.Draw(temp)
bbox = draw_temp.textbbox((0, 0), codigo_valor, font=font)
text_width = bbox[2] - bbox[0]
text_height = bbox[3] - bbox[1]

padding_text = 2
nuevo_alto = text_height + padding_text + alto_barcode_px
barcode_con_texto = Image.new("RGBA", (ancho_barcode_px, nuevo_alto), (255, 255, 255, 0))
draw = ImageDraw.Draw(barcode_con_texto)
draw.text(((ancho_barcode_px - text_width) // 2, 0), codigo_valor, fill="black", font=font)
barcode_con_texto.paste(barcode_img, (0, text_height + padding_text), barcode_img)

# --- 4. Crear fondo con imagen ---
fondo_img = Image.open("fondo.png").convert("RGBA")
fondo_img = fondo_img.resize((lado_qr_px + ancho_barcode_px + 20, max(lado_qr_px, nuevo_alto)))

# Pegar QR (izquierda)
pos_y_qr = (fondo_img.height - lado_qr_px) // 2
fondo_img.paste(qr_img, (0, pos_y_qr), qr_img)

# Pegar barcode (derecha)
pos_x_barcode = lado_qr_px + 20
pos_y_barcode = (fondo_img.height - nuevo_alto) // 2
fondo_img.paste(barcode_con_texto, (pos_x_barcode, pos_y_barcode), barcode_con_texto)

# --- 5. Guardar y mostrar ---
fondo_img.save("resultado_horizontal.png", dpi=(dpi, dpi))
fondo_img.show()
print("Imagen generada: resultado_horizontal.png")

# --- 6. Crear documento Word con etiquetas ---
doc = Document()

# Configuración de márgenes de la hoja
sections = doc.sections
for section in sections:
    section.top_margin = Cm(1)
    section.bottom_margin = Cm(1)
    section.left_margin = Cm(1)
    section.right_margin = Cm(1)

# Número de filas y columnas de etiquetas por página
filas = 5     # cantidad de filas (vertical)
columnas = 2  # cantidad de columnas (horizontal)

# Crear tabla para organizar las etiquetas
tabla = doc.add_table(rows=filas, cols=columnas)
tabla.autofit = False

for i in range(filas):
    for j in range(columnas):
        celda = tabla.cell(i, j)
        if os.path.exists("resultado_horizontal.png"):
            parrafo = celda.paragraphs[0]
            run = parrafo.add_run()
            # Ajusta el tamaño de la imagen dentro de cada celda
            run.add_picture("resultado_horizontal.png", width=Cm(7))
            parrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER

# Guardar el documento
doc.save("etiquetas.docx")
print("Documento Word generado: etiquetas.docx")
